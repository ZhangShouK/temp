import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter


BASE_DIR = Path(__file__).resolve().parent
INPUT_PATH = BASE_DIR / "pyrit_catalog.json"
OUTPUT_PATH = BASE_DIR / "pyrit_catalog.xlsx"


def add_header(sheet, headers: list[str]) -> None:
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(vertical="top", wrap_text=True)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions


def set_column_widths(sheet, widths: dict[int, float]) -> None:
    for column_index, width in widths.items():
        sheet.column_dimensions[get_column_letter(column_index)].width = width


def main() -> None:
    catalog = json.loads(INPUT_PATH.read_text(encoding="utf-8"))
    workbook = Workbook()

    scenarios_sheet = workbook.active
    scenarios_sheet.title = "Scenarios"
    add_header(scenarios_sheet, ["Scenario Name"])
    for scenario in sorted(catalog.get("scenarios", []), key=lambda item: item["registry_name"]):
        scenarios_sheet.append([scenario["registry_name"]])
    set_column_widths(scenarios_sheet, {1: 42})

    scenario_names_by_default_technique: dict[str, list[str]] = {}
    for scenario in catalog.get("scenarios", []):
        scenario_name = scenario["registry_name"]
        for technique_name in scenario.get("default_techniques", []):
            scenario_names_by_default_technique.setdefault(technique_name, []).append(scenario_name)

    techniques_sheet = workbook.create_sheet("Techniques")
    add_header(techniques_sheet, ["Technique Name", "Default Scenarios"])
    for technique in sorted(catalog.get("techniques", []), key=lambda item: item["name"]):
        scenario_names = sorted(set(scenario_names_by_default_technique.get(technique["name"], [])))
        techniques_sheet.append([technique["name"], "\n".join(scenario_names)])
    for row in techniques_sheet.iter_rows(min_row=2):
        row[1].alignment = Alignment(vertical="top", wrap_text=True)
    set_column_widths(techniques_sheet, {1: 36, 2: 48})

    datasets_sheet = workbook.create_sheet("Datasets")
    add_header(datasets_sheet, ["Dataset Name", "Harm Categories"])
    for dataset in sorted(catalog.get("datasets", []), key=lambda item: item["name"]):
        metadata = dataset.get("metadata") or {}
        harm_categories = metadata.get("harm_categories") or []
        datasets_sheet.append([dataset["name"], "\n".join(str(value) for value in harm_categories)])
    for row in datasets_sheet.iter_rows(min_row=2):
        row[1].alignment = Alignment(vertical="top", wrap_text=True)
    set_column_widths(datasets_sheet, {1: 46, 2: 32})

    workbook.save(OUTPUT_PATH)
    print(f"Wrote {OUTPUT_PATH} ({len(catalog.get('scenarios', []))} scenarios, "
          f"{len(catalog.get('techniques', []))} techniques, "
          f"{len(catalog.get('datasets', []))} datasets)")


if __name__ == "__main__":
    main()
